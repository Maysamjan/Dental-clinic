"""Excel export helper built on openpyxl."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def build_xlsx(title, header, rows):
    """Return an .xlsx file as bytes with a styled header row."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    ws.append(header)
    head_fill = PatternFill("solid", fgColor="0D9488")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill

    for row in rows:
        ws.append(row)

    # Auto-size columns roughly.
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
